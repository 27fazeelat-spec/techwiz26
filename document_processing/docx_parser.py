"""DOCX parsing with python-docx.

Keeps body order (paragraphs and tables interleaved), heading styles, list items,
table cells, footers, hidden (vanish-formatted) runs and core properties.
"""
import io
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from document_processing.models import Block, ParsedDocument

HEADING_STYLE = re.compile(r"^Heading (\d)$")


def _paragraph_block(paragraph, index, seen_title):
    text = paragraph.text
    if not text.strip():
        return None
    style = paragraph.style.name if paragraph.style is not None else ""
    hidden_parts = [r.text for r in paragraph.runs if r.font.hidden and r.text.strip()]
    block = Block(kind="paragraph", text=text, paragraph_index=index, style=style,
                  hidden=bool(hidden_parts), hidden_text=" ".join(hidden_parts))
    match = HEADING_STYLE.match(style)
    if style == "Title":
        block.kind = "title"
    elif match:
        block.kind, block.level = "heading", int(match.group(1))
    elif style.startswith("List"):
        block.kind = "list_item"
    elif not seen_title:
        block.kind = "letterhead"
    return block


def parse_docx(data):
    doc = Document(io.BytesIO(data))
    blocks, warnings = [], []
    index, seen_title = 0, False

    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            block = _paragraph_block(item, index, seen_title)
            index += 1
            if block:
                seen_title = seen_title or block.kind == "title"
                blocks.append(block)
        elif isinstance(item, Table):
            for row in item.rows:
                cells = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if not cells or cells[-1] != value:   # merged cells repeat their text
                        cells.append(value)
                if any(cells):
                    blocks.append(Block(kind="table_row", text=" | ".join(cells), cells=cells,
                                        paragraph_index=index))
            index += 1

    # If no Title style was used, the "letterhead" guess was wrong: treat those lines as body text.
    if not seen_title:
        for b in blocks:
            if b.kind == "letterhead":
                b.kind = "paragraph"

    seen_footer = set()
    for section in doc.sections:
        for p in section.footer.paragraphs:
            text = p.text.strip()
            if text and text not in seen_footer:
                seen_footer.add(text)
                hidden = [r.text for r in p.runs if r.font.hidden and r.text.strip()]
                blocks.append(Block(kind="footer", text=text, hidden=bool(hidden), hidden_text=" ".join(hidden)))

    props = doc.core_properties
    properties = {k: v for k, v in {
        "title": props.title, "subject": props.subject, "author": props.author,
        "keywords": props.keywords, "comments": props.comments,
    }.items() if v}

    if not any(b.kind == "heading" for b in blocks):
        warnings.append("No heading styles found; sections are identified by paragraph number.")
    return ParsedDocument(format="docx", blocks=blocks, properties=properties, warnings=warnings)

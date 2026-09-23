"""Reports and exports (SRS FR lxi-lxii, Steps 62-63, deliverables 1.10.6 and 1.10.8).

Every report is a query over stored runs, findings, progress and audit rows. Nothing is summarised
by GenAI and nothing is typed in by hand, so every number can be traced back to the row behind it.
A report is a Table (title, description, columns, rows); the same Table is rendered as HTML, CSV,
Excel or PDF.
"""
import csv
import io
from dataclasses import dataclass, field
from datetime import datetime

from sqlalchemy import case, func, select

from database import db, utcnow
from database.models import (AuditLog, ChangeImpact, ComparisonRow, Document, Employee, Finding, GenerationRun, JobRole,
                             MatrixRow, MatrixVersion, Plan, PlanItem, PlanItemRequirement, PlanModule, Progress,
                             QuizAttempt, Requirement, SecurityFinding, ValidationRun)

COMPARED_FIELDS = ["included", "mandatory", "priority", "due_stage", "module_category", "source_document", "source_section",
                   "task", "assessment_topic"]


@dataclass
class Table:
    name: str
    title: str
    description: str
    columns: list
    rows: list = field(default_factory=list)
    generated_at: datetime = field(default_factory=utcnow)
    pdf_columns: list | None = None        # a narrower set for the PDF page; CSV and Excel keep every column


def current_plans(role=None, employee_ids=None):
    query = select(Plan).where(Plan.status.notin_(["superseded", "generating"])).order_by(Plan.plan_code, Plan.version)
    if role:
        query = query.join(JobRole, Plan.job_role_id == JobRole.id).where(JobRole.code == role)
    if employee_ids is not None:
        query = query.where(Plan.employee_id.in_(employee_ids))
    return db.session.scalars(query).all()


def latest_runs(plans):
    runs = {}
    ids = [p.id for p in plans]
    if not ids:
        return runs
    for run in db.session.scalars(select(ValidationRun).where(ValidationRun.plan_id.in_(ids))
                                  .order_by(ValidationRun.created_at, ValidationRun.id)):
        runs[run.plan_id] = run                                # the last one wins
    return runs


def _pct(value):
    return None if value is None else round(value, 1)


# --------------------------------------------------------------------------- builders

def comparison(role=None, employee_ids=None):
    plans = current_plans(role, employee_ids)
    runs = latest_runs(plans)
    cols = ["Plan", "Employee", "Role", "Requirement", "Source"]
    for f in COMPARED_FIELDS:
        cols += [f"{f} (Python)", f"{f} (GenAI)", f"{f} match"]
    cols += ["All fields match", "Coverage", "Traceability", "Validation status", "Explanation"]
    t = Table("comparison", "GenAI vs Python comparison",
              "One row per requirement per plan: what the approved matrix (Python) expects next to what Gemini "
              "produced, field by field (SRS Step 46).", cols,
              pdf_columns=["Plan", "Requirement", "Source", "All fields match", "Coverage", "Validation status", "Explanation"])
    by_id = {p.id: p for p in plans}
    for plan_id, run in runs.items():
        p = by_id[plan_id]
        for r in db.session.scalars(select(ComparisonRow).where(ComparisonRow.validation_run_id == run.id)
                                    .order_by(ComparisonRow.req_id)):
            row = [f"{p.plan_code} v{p.version}", p.employee.name, r.role_code, r.req_id, r.source]
            for f in COMPARED_FIELDS:
                v = r.fields.get(f, {})
                row += [v.get("python", ""), v.get("genai", ""), "" if not v else ("Match" if v.get("match") else "Mismatch")]
            row += ["Yes" if r.overall_match else "No", r.coverage_status, r.traceability_status, r.validation_status, r.explanation]
            t.rows.append(row)
    return t


def validation(role=None, employee_ids=None):
    plans = current_plans(role, employee_ids)
    runs = latest_runs(plans)
    t = Table("validation", "Validation report",
              "Plan-level result of Pipeline 2 for every current plan: status, scores and finding counts (SRS 1.10.8).",
              ["Plan", "Employee", "Role", "Matrix", "Matrix approved", "Status", "Coverage %", "Traceability %",
               "Mandatory traceability %", "Requirement consistency %", "Generation consistency %", "Missing",
               "Unsupported", "Contradictions", "Duplicates", "Errors", "Warnings", "Ruleset", "Validated at (UTC)",
               "Assigned"])
    for p in plans:
        run = runs.get(p.id)
        sev = {}
        if run:
            sev = dict(db.session.execute(select(Finding.severity, func.count()).where(Finding.validation_run_id == run.id)
                                          .group_by(Finding.severity)).all())
        t.rows.append([f"{p.plan_code} v{p.version}", p.employee.name, p.job_role.code,
                       f"v{p.matrix_version.version_no}" if p.matrix_version else "", "Yes" if run and run.matrix_approved else "No",
                       p.status, _pct(p.score_coverage), _pct(p.score_traceability), _pct(p.score_traceability_mandatory),
                       _pct(p.score_requirement_consistency), _pct(p.score_generation_consistency), p.count_missing,
                       p.count_unsupported, p.count_contradictions, p.count_duplicates, sev.get("error", 0),
                       sev.get("warning", 0), run.ruleset_hash[:12] if run else "",
                       run.created_at.strftime("%Y-%m-%d %H:%M") if run else "", "Yes" if p.approved_at else "No"])
    return t


def findings(role=None, employee_ids=None):
    plans = current_plans(role, employee_ids)
    runs = latest_runs(plans)
    t = Table("findings", "Validation findings",
              "Every finding raised by the Python rules on current plans, with the rule that raised it.",
              ["Plan", "Rule", "Severity", "Implied status", "Item", "Requirement", "Message"])
    by_run = {r.id: p for p in plans for pid, r in runs.items() if pid == p.id}
    if by_run:
        for f in db.session.scalars(select(Finding).where(Finding.validation_run_id.in_(list(by_run)))
                                    .order_by(Finding.validation_run_id, Finding.severity, Finding.rule_id)):
            p = by_run[f.validation_run_id]
            t.rows.append([f"{p.plan_code} v{p.version}", f.rule_id, f.severity, f.status, f.item_key or "", f.req_id or "", f.message])
    return t


def hallucination(role=None, employee_ids=None):
    t = findings(role, employee_ids)
    t.name, t.title = "hallucination", "Hallucination and unsupported content"
    t.description = ("Findings where generated content states something the cited source does not support, cites no "
                     "valid source, or uses a requirement that does not apply (SRS Steps 31-32).")
    keep = {"V-HALLUCINATION", "V-SOURCE", "V-QUIZ-SOURCE", "V-REQ-ID", "V-ROLE", "V-CONDITION"}
    t.rows = [r for r in t.rows if r[1] in keep]
    return t


def traceability(role=None, employee_ids=None):
    plans = current_plans(role, employee_ids)
    t = Table("traceability", "Source traceability",
              "Every generated item with the requirements it is based on, the source it cites, whether that source "
              "exists in a current document, and the model and prompt that wrote it (SRS Step 15, challenge 7).",
              ["Plan", "Item", "Type", "Status", "Requirements", "Cited document", "Cited section", "Source found",
               "Model", "Prompt"])
    from database.models import Chunk
    current_sections = {tuple(r) for r in db.session.execute(select(Document.doc_id, Chunk.section_id).join(
        Chunk, Chunk.document_id == Document.id).where(Document.status.in_(["active", "expired"])))}
    runs = {r.id: r for r in db.session.scalars(select(GenerationRun).where(GenerationRun.plan_id.in_([p.id for p in plans])))} if plans else {}
    for p in plans:
        for m in p.modules:
            run = runs.get(m.run_id)
            for i in m.items:
                cited = (i.source_doc_id, i.source_section_id)
                t.rows.append([f"{p.plan_code} v{p.version}", i.item_key, i.item_type, i.status,
                               ", ".join(i.content.get("requirement_ids", [])), i.source_doc_id or "", i.source_section_id or "",
                               "" if not i.source_doc_id else ("Yes" if cited in current_sections else "No"),
                               run.model if run else "", f"{run.prompt_template} v{run.prompt_version}" if run else ""])
    return t


def progress(role=None, employee_ids=None):
    from config.settings import today
    from flask import current_app, has_app_context
    from src.services import progress as prog
    day = today(current_app.config) if has_app_context() else utcnow().date()
    query = select(Employee).order_by(Employee.employee_code)
    if employee_ids is not None:
        query = query.where(Employee.id.in_(employee_ids))
    t = Table("progress", "Employee progress",
              "Assigned plan, overall completion, progress status and completion by item type for every employee "
              "(SRS Steps 53-54). Status rules are in config/progress.yaml.",
              ["Employee", "Code", "Role", "Property", "Plan", "Overall %", "Status", "Overdue", "Checklist",
               "Quiz questions", "Tasks", "Assessments", "Waiting for sign-off"])
    for e in db.session.scalars(query):
        if role and e.job_role.code != role:
            continue
        plan = prog.assigned_plan(e)
        if plan is None:
            t.rows.append([e.name, e.employee_code, e.job_role.code, e.property.name, "Not assigned", "", "", "", "", "", "", "", ""])
            continue
        s = prog.summary(plan, e, day)
        part = lambda k: f"{s['by_type'][k]['done']}/{s['by_type'][k]['total']}" if k in s["by_type"] else ""
        t.rows.append([e.name, e.employee_code, e.job_role.code, e.property.name, f"{plan.plan_code} v{plan.version}",
                       s["pct"], s["status"], s["overdue"], part("checklist"), part("quiz_question"), part("task"),
                       part("assessment"), s["submitted"]])
    return t


def assessment_results(role=None, employee_ids=None):
    t = Table("assessment_results", "Quiz and assessment results",
              "Every quiz attempt (scored by Python against the validated answer key) and every assessment result "
              "recorded by a manager.",
              ["Employee", "Plan", "Kind", "Module", "Attempt", "Score", "Result", "Recorded by", "When (UTC)"])
    query = select(QuizAttempt, Employee, Plan).join(Employee, QuizAttempt.employee_id == Employee.id).join(
        Plan, QuizAttempt.plan_id == Plan.id).order_by(Employee.employee_code, QuizAttempt.submitted_at)
    for a, e, p in db.session.execute(query):
        if (role and p.job_role.code != role) or (employee_ids is not None and e.id not in employee_ids):
            continue
        t.rows.append([e.name, f"{p.plan_code} v{p.version}", "Quiz", a.module_key, a.attempt_no, a.score,
                       "Passed" if a.passed else "Not passed", "Python", a.submitted_at.strftime("%Y-%m-%d %H:%M")])
    query = (select(Progress, PlanItem, Employee, Plan).join(PlanItem, Progress.plan_item_id == PlanItem.id)
             .join(Employee, Progress.employee_id == Employee.id).join(Plan, Progress.plan_id == Plan.id)
             .where(PlanItem.item_type == "assessment", Progress.score.is_not(None)))
    for r, i, e, p in db.session.execute(query):
        if (role and p.job_role.code != role) or (employee_ids is not None and e.id not in employee_ids):
            continue
        t.rows.append([e.name, f"{p.plan_code} v{p.version}", "Assessment", i.module.module_key, "", r.score,
                       "Passed" if r.status == "completed" else "Not passed", r.verified_by or "",
                       r.completed_at.strftime("%Y-%m-%d %H:%M") if r.completed_at else ""])
    return t


def role_coverage(role=None, employee_ids=None):
    matrix = db.session.scalar(select(MatrixVersion).where(MatrixVersion.status == "approved").order_by(MatrixVersion.version_no.desc()))
    t = Table("role_coverage", "Role coverage",
              f"Per role: requirements in the approved matrix (v{matrix.version_no if matrix else '-'}), employees, "
              "current plans and their average mandatory coverage.",
              ["Role", "Name", "Status", "Matrix requirements", "Mandatory", "Employees", "Current plans",
               "Verified plans", "Average coverage %", "Assigned plans"])
    plans = current_plans()
    for r in db.session.scalars(select(JobRole).order_by(JobRole.code)):
        if role and r.code != role:
            continue
        total = mand = 0
        if matrix:
            total, mand = db.session.execute(select(func.count(), func.coalesce(func.sum(
                case((MatrixRow.mandatory.is_(True), 1), else_=0)), 0)).where(MatrixRow.matrix_version_id == matrix.id,
                                                                              MatrixRow.job_role_id == r.id)).one()
        mine = [p for p in plans if p.job_role_id == r.id]
        cov = [p.score_coverage for p in mine if p.score_coverage is not None]
        staff = db.session.scalar(select(func.count()).select_from(Employee).where(Employee.job_role_id == r.id))
        t.rows.append([r.code, r.name, r.status, total, mand, staff, len(mine),
                       sum(1 for p in mine if p.status in ("Verified", "Verified with Warning")),
                       round(sum(cov) / len(cov), 1) if cov else "", sum(1 for p in mine if p.approved_at)])
    return t


def mandatory_training(role=None, employee_ids=None):
    plans = current_plans(role, employee_ids)
    runs = latest_runs(plans)
    t = Table("mandatory_training", "Mandatory training coverage",
              "For every current plan, each mandatory requirement the approved matrix expects and whether the plan "
              "covers it (the numerator and denominator of the coverage score).",
              ["Plan", "Employee", "Requirement", "Source", "Covered", "Validation status"])
    by_id = {p.id: p for p in plans}
    for plan_id, run in runs.items():
        p = by_id[plan_id]
        for r in db.session.scalars(select(ComparisonRow).where(ComparisonRow.validation_run_id == run.id)
                                    .order_by(ComparisonRow.req_id)):
            if r.coverage_status in ("Covered", "Missing"):         # mandatory in the approved matrix
                t.rows.append([f"{p.plan_code} v{p.version}", p.employee.name, r.req_id, r.source,
                               "Yes" if r.coverage_status == "Covered" else "No", r.validation_status])
    return t


def policy_coverage(role=None, employee_ids=None):
    t = Table("policy_coverage", "Policy coverage",
              "Per current source document: requirements extracted, how many are mandatory, how many reach the "
              "approved matrix, how many are used by current plan items, and recorded version changes.",
              ["Document", "Version", "Title", "Category", "Status", "Requirements", "Mandatory", "In approved matrix",
               "Used by plan items", "Version changes"])
    matrix = db.session.scalar(select(MatrixVersion).where(MatrixVersion.status == "approved").order_by(MatrixVersion.version_no.desc()))
    in_matrix = set(db.session.scalars(select(MatrixRow.requirement_id).where(MatrixRow.matrix_version_id == matrix.id))) if matrix else set()
    used = set(db.session.scalars(select(PlanItemRequirement.requirement_id).join(PlanItem).join(PlanModule).join(Plan)
                                  .where(Plan.status.notin_(["superseded", "generating"]))))
    changes = dict(db.session.execute(select(ChangeImpact.doc_id, func.count()).group_by(ChangeImpact.doc_id)).all())
    for d in db.session.scalars(select(Document).where(Document.status.in_(["active", "expired"]), Document.tier > 0)
                                .order_by(Document.doc_id)):
        reqs = db.session.scalars(select(Requirement).where(Requirement.document_id == d.id)).all()
        t.rows.append([d.doc_id, d.version, d.title, d.category, d.status, len(reqs), sum(r.mandatory for r in reqs),
                       sum(r.id in in_matrix for r in reqs), sum(r.id in used for r in reqs), changes.get(d.doc_id, 0)])
    return t


def security(role=None, employee_ids=None):
    t = Table("security", "Security findings",
              "Prompt-injection and adversarial content caught at ingestion (SRS Steps 42-43, security testing report).",
              ["Document", "Version", "Section", "Technique", "Pattern", "Severity", "Action", "Excerpt"])
    from database.models import Chunk
    for f, d, c in db.session.execute(select(SecurityFinding, Document, Chunk).join(Document, SecurityFinding.document_id == Document.id)
                                      .outerjoin(Chunk, SecurityFinding.chunk_id == Chunk.id).order_by(Document.doc_id)):
        t.rows.append([d.doc_id, d.version, c.section_id if c else "whole document", f.technique, f.pattern or "",
                       f.severity, f.action, f.excerpt])
    return t


def audit_trail(role=None, employee_ids=None, limit=2000):
    t = Table("audit", "Audit trail",
              f"The most recent {limit} entries of the append-only audit log.",
              ["When (UTC)", "Action", "Entity", "ID", "Version", "By", "Reason"])
    for e in db.session.scalars(select(AuditLog).order_by(AuditLog.ts.desc()).limit(limit)):
        t.rows.append([e.ts.strftime("%Y-%m-%d %H:%M:%S"), e.action, e.entity_type, e.entity_id, e.entity_version or "",
                       (e.actor or {}).get("email") or (e.actor or {}).get("user_id", ""), e.reason or ""])
    return t


REPORTS = {
    "comparison": comparison, "validation": validation, "findings": findings, "hallucination": hallucination,
    "traceability": traceability, "mandatory_training": mandatory_training, "role_coverage": role_coverage,
    "policy_coverage": policy_coverage, "progress": progress, "assessment_results": assessment_results,
    "security": security, "audit": audit_trail,
}
TEAM_REPORTS = ("progress", "assessment_results")          # what a line manager may export for their team


def build(name, **filters):
    if name not in REPORTS:
        raise KeyError(name)
    return REPORTS[name](**filters)


# --------------------------------------------------------------------------- renderers

def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return v


def to_csv(t):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(t.columns)
    for r in t.rows:
        w.writerow([_cell(v) for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")        # BOM so Excel opens UTF-8 correctly


def to_xlsx(t):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = t.name[:31]
    ws.append(t.columns)
    for r in t.rows:
        ws.append([_cell(v) if not isinstance(v, (list, dict)) else str(v) for v in r])
    head = PatternFill("solid", fgColor="E3F4EC")
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True), head, Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if t.rows:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(t.columns, start=1):
        width = max([len(str(col))] + [len(str(_cell(r[i - 1]))) for r in t.rows[:300]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    about = wb.create_sheet("About")
    for line in [[t.title], [t.description], [f"Rows: {len(t.rows)}"],
                 [f"Generated (UTC): {t.generated_at:%Y-%m-%d %H:%M}"], ["Source: SkillSprint AI, computed from stored data"]]:
        about.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(t, max_rows=1500):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=6.5, leading=8)
    headcell = cell.clone("head", fontName="Helvetica-Bold")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=t.title, author="SkillSprint AI")
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    story = [Paragraph(esc(t.title), styles["Title"]), Paragraph(esc(t.description), styles["BodyText"]),
             Paragraph(f"{len(t.rows)} rows · generated {t.generated_at:%Y-%m-%d %H:%M} UTC · computed from stored data",
                       styles["Italic"]), Spacer(1, 4 * mm)]
    columns = t.pdf_columns or t.columns[:14]
    keep = [t.columns.index(c) for c in columns]
    rows = [[r[i] for i in keep] for r in t.rows[:max_rows]]
    if len(columns) < len(t.columns):
        story.append(Paragraph(f"The PDF shows {len(columns)} of {len(t.columns)} columns; the CSV and Excel exports contain all of them.",
                               styles["Italic"]))
    data = [[Paragraph(esc(c), headcell) for c in columns]] + [[Paragraph(esc(_cell(v)), cell) for v in r] for r in rows]
    usable = landscape(A4)[0] - 20 * mm
    weights = [max(6, min(40, max([len(str(c))] + [len(str(_cell(r[i]))) for r in rows[:200]]))) for i, c in enumerate(columns)]
    widths = [usable * w / sum(weights) for w in weights]
    table = LongTable(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3F4EC")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CFD7D3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9F8")]),
    ]))
    story.append(table)
    if len(t.rows) > max_rows:
        story.append(Paragraph(f"Showing the first {max_rows} of {len(t.rows)} rows; the CSV and Excel exports contain all rows.",
                               styles["Italic"]))
    doc.build(story)
    return buf.getvalue()


FORMATS = {"csv": ("text/csv; charset=utf-8", to_csv), "xlsx": (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", to_xlsx), "pdf": ("application/pdf", to_pdf)}

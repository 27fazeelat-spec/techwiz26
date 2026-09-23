"""Build the Aurelle sample document collection (PDF + DOCX) from text sources.

Each file in sources/ is one document: YAML front matter + a lightly marked-up body.
Every entry under `versions:` produces one output file in sample_documents/.

Body markup
  ## 1. Heading            level-1 heading
  ### 2.1 Heading          level-2 heading (the section id is the leading number)
  - item                   bullet
  | a | b |                table row (first row = header; |---| rows are skipped)
  blank line               paragraph break
  **bold**                 bold
  {{hidden:text}}          hidden text (DOCX vanish / PDF white 1pt)
  {{zw:text}}              zero-width spaces inserted between the letters
  {{b64:text}}             replaced by the base64 encoding of text

Version entries may carry `replace: [[old, new], ...]` and `remove: [section ids]`
to derive older or newer versions from the same source.

Usage:  python tools/dataset_builder/build_documents.py
"""
import base64
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parents[2]
SOURCES = Path(__file__).parent / "sources"
OUT = ROOT / "sample_documents"

ZWSP = "\u200b"
INK = "#1F2328"
MUTED = "#5B616B"
BRAND = "#7A2E3A"   # Aurelle's own letterhead colour (the client brand, not SkillSprint's)


# --------------------------------------------------------------------------- parsing

def load_source(path):
    text = path.read_text(encoding="utf-8")
    _, front, body = text.split("---", 2)
    return yaml.safe_load(front), body.strip("\n")


def apply_version(body, version):
    for old, new in version.get("replace", []):
        if old not in body:
            raise ValueError(f"replace target not found: {old[:60]!r}")
        body = body.replace(old, new)
    for sec in version.get("remove", []):
        body = remove_section(body, sec)
    return body


def remove_section(body, sec_id):
    lines, out, skipping, level = body.split("\n"), [], False, 0
    for line in lines:
        m = re.match(r"^(#{2,3}) (\S+)", line)
        if m:
            this_level = len(m.group(1))
            if m.group(2).rstrip(".") == sec_id:
                skipping, level = True, this_level
                continue
            if skipping and this_level <= level:
                skipping = False
        if not skipping:
            out.append(line)
    return "\n".join(out)


def parse_blocks(body):
    """Return a list of (kind, payload) blocks."""
    blocks, para, bullets, table = [], [], [], []

    def flush():
        nonlocal para, bullets, table
        if para:
            blocks.append(("p", " ".join(para)))
        if bullets:
            blocks.append(("ul", bullets))
        if table:
            blocks.append(("table", table))
        para, bullets, table = [], [], []

    for raw in body.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            flush()
        elif line.startswith("### "):
            flush(); blocks.append(("h2", line[4:].strip()))
        elif line.startswith("## "):
            flush(); blocks.append(("h1", line[3:].strip()))
        elif line.startswith("- "):
            if para or table:
                flush()
            bullets.append(line[2:].strip())
        elif line.startswith("|"):
            if para or bullets:
                flush()
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                table.append(cells)
        else:
            if bullets or table:
                flush()
            para.append(line.strip())
    flush()
    return blocks


def expand_inline(text):
    """Split text into runs: list of (text, style) where style in {'', 'b', 'hidden'}."""
    text = re.sub(r"\{\{b64:(.*?)\}\}", lambda m: base64.b64encode(m.group(1).encode()).decode(), text)
    text = re.sub(r"\{\{zw:(.*?)\}\}", lambda m: ZWSP.join(m.group(1)), text)
    runs, pos = [], 0
    for m in re.finditer(r"\{\{hidden:(.*?)\}\}|\*\*(.+?)\*\*", text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], ""))
        runs.append((m.group(1), "hidden") if m.group(1) is not None else (m.group(2), "b"))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], ""))
    return runs


def metadata_rows(meta, ver):
    return [
        ("Document ID", meta["doc_id"]),
        ("Title", ver.get("title", meta["title"])),
        ("Version", ver["version"]),
        ("Status", ver.get("status", "Approved")),
        ("Effective Date", ver.get("effective", "")),
        ("Review / Expiry Date", ver.get("review", "")),
        ("Owner Department", meta["owner"]),
        ("Category", meta["category"]),
        ("Applies To", meta["applies_to"]),
        ("Supersedes", ver.get("supersedes", "None")),
    ]


def revision_rows(meta, ver):
    rows = meta.get("revision_history", [])
    upto = ver.get("revision_upto", ver["version"])
    order = [str(r["version"]) for r in rows]
    if str(upto) in order:
        rows = rows[: order.index(str(upto)) + 1]
    return [(str(r["version"]), str(r["date"]), r["summary"]) for r in rows]


# --------------------------------------------------------------------------- DOCX

def docx_add_runs(paragraph, text):
    for chunk, style in expand_inline(text):
        run = paragraph.add_run(chunk)
        if style == "b":
            run.bold = True
        elif style == "hidden":
            run.font.hidden = True


def docx_page_field(paragraph):
    run = paragraph.add_run()
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)


def docx_table(doc, rows, header=True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            docx_add_runs(cell.paragraphs[0], value)
            if header and r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()
    return table


def build_docx(meta, ver, blocks, path):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    letter = doc.add_paragraph()
    run = letter.add_run(meta.get("letterhead", "AURELLE HOTELS & RESIDENCES"))
    run.bold, run.font.size = True, Pt(10)
    run.font.color.rgb = RGBColor.from_string(BRAND.lstrip("#"))

    if ver.get("watermark"):
        mark = doc.add_paragraph()
        wrun = mark.add_run(ver["watermark"])
        wrun.bold, wrun.font.size = True, Pt(14)
        wrun.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)

    doc.add_paragraph(ver.get("title", meta["title"]), style="Title")

    if meta.get("header", True):
        rows = metadata_rows(meta, ver)
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        for i, (k, v) in enumerate(rows):
            table.cell(i, 0).text = k
            table.cell(i, 1).text = str(v)
            table.cell(i, 0).paragraphs[0].runs[0].bold = True
        doc.add_paragraph()

    for kind, payload in blocks:
        if kind == "h1":
            doc.add_heading(payload, level=1)
        elif kind == "h2":
            doc.add_heading(payload, level=2)
        elif kind == "p":
            docx_add_runs(doc.add_paragraph(), payload)
        elif kind == "ul":
            for item in payload:
                docx_add_runs(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "table":
            docx_table(doc, payload)

    history = revision_rows(meta, ver)
    if history and meta.get("header", True):
        doc.add_heading("Revision History", level=1)
        docx_table(doc, [("Version", "Date", "Summary of changes")] + history)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    label = f"{meta['doc_id']} v{ver['version']} · {ver.get('title', meta['title'])} · Page "
    footer.add_run(label).font.size = Pt(8)
    docx_page_field(footer)
    if meta.get("footer_note"):
        extra = doc.sections[0].footer.add_paragraph()
        extra.add_run(meta["footer_note"]).font.size = Pt(6)

    props = doc.core_properties
    props.title = ver.get("title", meta["title"])
    props.author = meta.get("author", "Aurelle Hospitality Group")
    props.subject = f"{meta['doc_id']} v{ver['version']}"
    doc.save(path)


# --------------------------------------------------------------------------- PDF

def register_fonts():
    fonts = Path("C:/Windows/Fonts")
    if (fonts / "arial.ttf").exists():
        pdfmetrics.registerFont(TTFont("Body", str(fonts / "arial.ttf")))
        pdfmetrics.registerFont(TTFont("Body-Bold", str(fonts / "arialbd.ttf")))
        pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-Bold")
        return "Body", "Body-Bold"
    return "Helvetica", "Helvetica-Bold"


FONT, FONT_B = register_fonts()
S = {
    "letter": ParagraphStyle("letter", fontName=FONT_B, fontSize=9, textColor=BRAND, spaceAfter=10),
    "title": ParagraphStyle("title", fontName=FONT_B, fontSize=20, leading=24, textColor=INK, spaceAfter=12),
    "h1": ParagraphStyle("h1", fontName=FONT_B, fontSize=13.5, leading=17, textColor=INK, spaceBefore=12, spaceAfter=6),
    "h2": ParagraphStyle("h2", fontName=FONT_B, fontSize=11, leading=14, textColor=INK, spaceBefore=8, spaceAfter=4),
    "p": ParagraphStyle("p", fontName=FONT, fontSize=10, leading=14, textColor=INK, spaceAfter=6, alignment=TA_LEFT),
    "cell": ParagraphStyle("cell", fontName=FONT, fontSize=9, leading=12, textColor=INK),
    "cellb": ParagraphStyle("cellb", fontName=FONT_B, fontSize=9, leading=12, textColor=INK),
    "hidden": ParagraphStyle("hidden", fontName=FONT, fontSize=1, leading=2, textColor=colors.white),
    "warn": ParagraphStyle("warn", fontName=FONT_B, fontSize=14, textColor="#B03030", spaceAfter=8),
}


def esc(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf_markup(text):
    """Return (visible markup, hidden texts)."""
    out, hidden = [], []
    for chunk, style in expand_inline(text):
        if style == "hidden":
            hidden.append(chunk)
        elif style == "b":
            out.append(f"<b>{esc(chunk)}</b>")
        else:
            out.append(esc(chunk))
    return "".join(out), hidden


def pdf_paragraphs(text, style):
    visible, hidden = pdf_markup(text)
    flow = [Paragraph(visible, style)] if visible.strip() else []
    flow += [Paragraph(esc(h), S["hidden"]) for h in hidden]
    return flow


def pdf_table(rows, widths=None, header=True):
    data = []
    for r, row in enumerate(rows):
        style = S["cellb"] if header and r == 0 else S["cell"]
        data.append([Paragraph(pdf_markup(c)[0], style) for c in row])
    table = Table(data, colWidths=widths, hAlign="LEFT", repeatRows=1 if header else 0)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C9CDD2")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F2F4") if header else colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def build_pdf(meta, ver, blocks, path):
    title = ver.get("title", meta["title"])
    footer_text = f"{meta['doc_id']} v{ver['version']} · {title}"

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(FONT, 7.5)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(20 * mm, 12 * mm, f"{footer_text} · Page {doc.page}")
        if ver.get("watermark"):
            canvas.setFont(FONT_B, 54)
            canvas.setFillColor(colors.Color(0.75, 0.2, 0.2, alpha=0.12))
            canvas.translate(A4[0] / 2, A4[1] / 2)
            canvas.rotate(35)
            canvas.drawCentredString(0, 0, "DRAFT - NOT APPROVED")
        canvas.restoreState()

    story = [Paragraph(esc(meta.get("letterhead", "AURELLE HOTELS & RESIDENCES")), S["letter"])]
    if ver.get("watermark"):
        story.append(Paragraph(esc(ver["watermark"]), S["warn"]))
    story.append(Paragraph(esc(title), S["title"]))
    if meta.get("header", True):
        rows = [(k, str(v)) for k, v in metadata_rows(meta, ver)]
        story += [pdf_table(rows, widths=[45 * mm, 120 * mm], header=False), Spacer(1, 8)]

    for kind, payload in blocks:
        if kind in ("h1", "h2"):
            story.append(Paragraph(esc(payload), S[kind]))
        elif kind == "p":
            story += pdf_paragraphs(payload, S["p"])
        elif kind == "ul":
            items = [ListItem(Paragraph(pdf_markup(i)[0], S["p"]), leftIndent=12) for i in payload]
            story.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=12))
        elif kind == "table":
            story += [pdf_table(payload), Spacer(1, 6)]

    history = revision_rows(meta, ver)
    if history and meta.get("header", True):
        story.append(Paragraph("Revision History", S["h1"]))
        story.append(pdf_table([("Version", "Date", "Summary of changes")] + history,
                               widths=[22 * mm, 28 * mm, 115 * mm]))

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=20 * mm, title=title,
                            author=meta.get("author", "Aurelle Hospitality Group"),
                            subject=f"{meta['doc_id']} v{ver['version']}")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


# --------------------------------------------------------------------------- main

def main():
    OUT.mkdir(exist_ok=True)
    only = set(sys.argv[1:])
    built = 0
    for src in sorted(SOURCES.glob("*.md")):
        meta, body = load_source(src)
        if only and meta["doc_id"] not in only:
            continue
        for ver in meta["versions"]:
            blocks = parse_blocks(apply_version(body, ver))
            ext = meta["format"].lower()
            name = f"{meta['doc_id']}_v{ver['version']}.{ext}"
            (build_pdf if ext == "pdf" else build_docx)(meta, ver, blocks, OUT / name)
            print(f"  built {name}")
            built += 1
    print(f"{built} file(s) written to {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
